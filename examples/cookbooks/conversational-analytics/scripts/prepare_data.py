"""Download the original UCI workbook and build a reproducible, local sales database."""

import hashlib
import json
import sqlite3
from collections import Counter
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from urllib.request import urlopen
from zipfile import ZipFile

from openpyxl import load_workbook

SOURCE = "https://archive.ics.uci.edu/static/public/352/online%2Bretail.zip"
SHA256 = "f5385cbb54bbebf7196389109c6b0621faab0c304e3702548165e71c84aede8b"
DATA = Path(__file__).resolve().parents[1] / "data"


def sales_row(row):
    invoice, stock, description, quantity, date, price, _customer, country = row
    if str(invoice).upper().startswith("C"):
        return None, "cancellation"
    if quantity <= 0:
        return None, "non_positive_quantity"
    if price <= 0:
        return None, "non_positive_price"
    # Retain fractional-penny unit prices exactly. Round only for display.
    units = Decimal(str(price)) * 10000 * quantity
    if units != units.to_integral_value() or not isinstance(date, datetime):
        raise ValueError("Unexpected source price precision or invoice date")
    return (str(invoice), str(stock), str(description or stock).strip(), int(quantity),
            date.date().isoformat(), str(country), int(units)), None


def main():
    DATA.mkdir(exist_ok=True)
    archive = DATA / "online-retail.zip"
    if not archive.exists():
        print("Downloading UCI Online Retail (about 23 MB)...", flush=True)
        temporary = archive.with_suffix(".download")
        with urlopen(SOURCE, timeout=120) as response, temporary.open("wb") as output:
            while chunk := response.read(1024 * 1024):
                output.write(chunk)
        temporary.replace(archive)
    digest = hashlib.sha256(archive.read_bytes()).hexdigest()
    if digest != SHA256:
        raise ValueError("Source checksum changed. Remove the cached archive and verify the UCI release before importing.")
    temporary_db = DATA / "retail.build.sqlite"
    temporary_db.unlink(missing_ok=True)
    db = sqlite3.connect(temporary_db)
    db.execute("CREATE TABLE sales (invoice TEXT, stock_code TEXT, description TEXT, quantity INTEGER, day TEXT, country TEXT, sales_units INTEGER)")
    counts = Counter()
    with ZipFile(archive) as bundle, bundle.open("Online Retail.xlsx") as workbook_file:
        workbook = load_workbook(workbook_file, read_only=True, data_only=True)
        rows = workbook.active.iter_rows(values_only=True)
        assert next(rows) == ("InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate", "UnitPrice", "CustomerID", "Country")
        batch = []
        for row in rows:
            counts["source_rows"] += 1
            record, excluded = sales_row(row)
            if record is None:
                counts[excluded] += 1
                continue
            counts["eligible_rows"] += 1
            batch.append(record)
            if len(batch) == 5000:
                db.executemany("INSERT INTO sales VALUES (?, ?, ?, ?, ?, ?, ?)", batch)
                batch.clear()
        db.executemany("INSERT INTO sales VALUES (?, ?, ?, ?, ?, ?, ?)", batch)
        workbook.close()
    if counts["source_rows"] != 541909:
        raise ValueError("Unexpected source row count")
    db.executescript("CREATE INDEX sales_day ON sales(day); CREATE INDEX sales_country_day ON sales(country, day);")
    db.commit()
    db.close()
    temporary_db.replace(DATA / "retail.sqlite")
    manifest = {"source": SOURCE, "sha256": digest, "license": "CC BY 4.0", **counts}
    (DATA / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    print(json.dumps(manifest, indent=2))
    print("Ready: data/retail.sqlite")


if __name__ == "__main__":
    main()
